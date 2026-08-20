"""Tests for /exportar refeicoes: format + period parsing, CSV/JSON/XLSX export."""

from __future__ import annotations

import json
import os
import tempfile
from datetime import date, timedelta
from unittest.mock import AsyncMock, MagicMock, patch

import pytest

from src.database.repository import Repository
from src.config import Config
from src.telegram.bot import TelegramBot
from src.telegram.helpers import _parse_export_args


# ---------------------------------------------------------------------------
# _parse_export_args
# ---------------------------------------------------------------------------

class TestParseExportArgs:
    def test_defaults_to_csv_last_90_days(self):
        today = date.today()
        fmt, start, end = _parse_export_args([])
        assert fmt == "csv"
        assert end == today
        assert start == today - timedelta(days=90)

    def test_format_only(self):
        fmt, start, end = _parse_export_args(["json"])
        assert fmt == "json"

    def test_format_is_case_insensitive(self):
        fmt, _, _ = _parse_export_args(["XLSX"])
        assert fmt == "xlsx"

    def test_format_and_start_date(self):
        fmt, start, end = _parse_export_args(["xlsx", "2026-08-01"])
        assert fmt == "xlsx"
        assert start == date(2026, 8, 1)
        assert end == date.today()

    def test_format_and_full_period(self):
        fmt, start, end = _parse_export_args(["csv", "2026-08-01", "2026-08-10"])
        assert (fmt, start, end) == ("csv", date(2026, 8, 1), date(2026, 8, 10))

    def test_dates_without_format_default_csv(self):
        fmt, start, end = _parse_export_args(["2026-08-01", "2026-08-10"])
        assert fmt == "csv"
        assert (start, end) == (date(2026, 8, 1), date(2026, 8, 10))

    def test_invalid_date_raises(self):
        with pytest.raises(ValueError, match="Data inválida"):
            _parse_export_args(["not-a-date"])

    def test_start_after_end_raises(self):
        with pytest.raises(ValueError, match="início"):
            _parse_export_args(["2026-08-10", "2026-08-01"])

    def test_future_start_raises(self):
        future = (date.today() + timedelta(days=5)).isoformat()
        with pytest.raises(ValueError, match="futuras"):
            _parse_export_args([future])

    def test_too_many_args_raises(self):
        with pytest.raises(ValueError, match="Demasiados"):
            _parse_export_args(["csv", "2026-08-01", "2026-08-10", "extra"])


# ---------------------------------------------------------------------------
# /exportar refeicoes handler
# ---------------------------------------------------------------------------

def _make_config(**overrides):
    defaults = {
        "telegram_bot_token": "fake-token",
        "telegram_chat_id": "123456",
        "garmin_email": "test@example.com",
        "garmin_password": "secret",
        "database_path": ":memory:",
        "log_level": "INFO",
        "log_file": None,
        "daily_alerts": False,
        "groq_api_key": None,
        "usda_api_key": None,
        "api_ninjas_key": None,
        "fatsecret_consumer_key": None,
        "fatsecret_consumer_secret": None,
        "garmin_api_port": None,
        "garmin_api_key": None,
        "health_port": None,
        "newsletter_enabled": False,
    }
    defaults.update(overrides)
    cfg = MagicMock(spec=Config)
    for k, v in defaults.items():
        setattr(cfg, k, v)
    return cfg


@pytest.fixture
def db_path():
    with tempfile.NamedTemporaryFile(suffix=".db", delete=False) as f:
        path = f.name
    yield path
    try:
        os.unlink(path)
    except PermissionError:
        pass


@pytest.fixture
def repo(db_path):
    r = Repository(db_path)
    r.init_database()
    yield r
    r._engine.dispose()


_NEXT_CHAT_ID = 800000


def _make_update():
    global _NEXT_CHAT_ID
    _NEXT_CHAT_ID += 1
    update = MagicMock()
    update.effective_chat.id = _NEXT_CHAT_ID
    update.message = AsyncMock()
    update.message.reply_text = AsyncMock()
    return update


def _make_bot(repo, chat_id, fatsecret_client=None):
    cfg = _make_config(telegram_chat_id=str(chat_id))
    bot = TelegramBot(cfg, repo, fatsecret_client=fatsecret_client)
    bot._chat_id = chat_id
    return bot


def _make_context(args):
    ctx = MagicMock()
    ctx.args = args
    return ctx


class TestCmdExportarRefeicoes:
    @pytest.mark.asyncio
    async def test_no_entries_replies_with_message(self, repo):
        update = _make_update()
        bot = _make_bot(repo, update.effective_chat.id)
        await bot._cmd_exportar(update, _make_context(["refeicoes"]))
        update.message.reply_text.assert_awaited_once()
        assert "Sem dados" in update.message.reply_text.call_args[0][0]

    @pytest.mark.asyncio
    async def test_invalid_period_replies_with_error(self, repo):
        update = _make_update()
        bot = _make_bot(repo, update.effective_chat.id)
        await bot._cmd_exportar(update, _make_context(["refeicoes", "csv", "2026-08-10", "2026-08-01"]))
        update.message.reply_text.assert_awaited_once()
        assert "❌" in update.message.reply_text.call_args[0][0]

    @pytest.mark.asyncio
    async def test_csv_export_sends_document(self, repo):
        today = date.today()
        repo.save_food_entries(today, [
            {"name": "Ovos", "quantity": 2, "unit": "un", "calories": 140.0,
             "protein_g": 12.0, "fat_g": 10.0, "carbs_g": 1.0, "fiber_g": 0.0,
             "source": "llm_estimate"},
        ])
        update = _make_update()
        bot = _make_bot(repo, update.effective_chat.id)

        with patch("telegram.Bot.send_document", new_callable=AsyncMock) as mock_send:
            await bot._cmd_exportar(update, _make_context(["refeicoes"]))

        mock_send.assert_awaited_once()
        kwargs = mock_send.call_args.kwargs
        assert kwargs["chat_id"] == bot._chat_id
        assert kwargs["document"].filename.endswith(".csv")
        content = kwargs["document"].input_file_content
        assert b"Ovos" in content

    @pytest.mark.asyncio
    async def test_json_export_contains_valid_json(self, repo):
        today = date.today()
        repo.save_food_entries(today, [
            {"name": "Frango", "quantity": 150, "unit": "g", "calories": 250.0,
             "protein_g": 40.0, "fat_g": 6.0, "carbs_g": 0.0, "fiber_g": 0.0,
             "source": "fatsecret", "barcode": "FS1"},
        ])
        update = _make_update()
        bot = _make_bot(repo, update.effective_chat.id)

        with patch("telegram.Bot.send_document", new_callable=AsyncMock) as mock_send:
            await bot._cmd_exportar(update, _make_context(["refeicoes", "json"]))

        content = mock_send.call_args.kwargs["document"].input_file_content
        data = json.loads(content)
        assert data[0]["nome"] == "Frango"
        assert data[0]["fonte"] == "fatsecret"

    @pytest.mark.asyncio
    async def test_xlsx_export_is_valid_workbook(self, repo):
        today = date.today()
        repo.save_food_entries(today, [
            {"name": "Iogurte", "quantity": 1, "unit": "un", "calories": 90.0,
             "protein_g": 8.0, "fat_g": 2.0, "carbs_g": 10.0, "fiber_g": 0.0,
             "source": "openfoodfacts"},
        ])
        update = _make_update()
        bot = _make_bot(repo, update.effective_chat.id)

        with patch("telegram.Bot.send_document", new_callable=AsyncMock) as mock_send:
            await bot._cmd_exportar(update, _make_context(["refeicoes", "xlsx"]))

        content = mock_send.call_args.kwargs["document"].input_file_content
        import io
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(content))
        rows = list(wb.active.iter_rows(values_only=True))
        assert rows[1][1] == "Iogurte"

    @pytest.mark.asyncio
    async def test_custom_period_passed_to_repository(self, repo):
        update = _make_update()
        bot = _make_bot(repo, update.effective_chat.id)

        with patch.object(repo, "get_food_entries_range", wraps=repo.get_food_entries_range) as mock_range:
            await bot._cmd_exportar(
                update, _make_context(["refeicoes", "csv", "2026-08-01", "2026-08-10"])
            )

        mock_range.assert_called_once_with(date(2026, 8, 1), date(2026, 8, 10))

    @pytest.mark.asyncio
    async def test_nutricao_alias_still_works(self, repo):
        """Backward compatibility: /exportar nutricao behaves like /exportar refeicoes."""
        update = _make_update()
        bot = _make_bot(repo, update.effective_chat.id)
        await bot._cmd_exportar(update, _make_context(["nutricao"]))
        update.message.reply_text.assert_awaited_once()
        assert "Sem dados" in update.message.reply_text.call_args[0][0]


# ---------------------------------------------------------------------------
# FatSecret pre-sync before export
# ---------------------------------------------------------------------------

class TestSyncMissingFatsecretDays:
    @pytest.mark.asyncio
    async def test_no_fatsecret_client_skips_sync_silently(self, repo):
        update = _make_update()
        bot = _make_bot(repo, update.effective_chat.id, fatsecret_client=None)

        with patch("telegram.Bot.send_document", new_callable=AsyncMock):
            await bot._cmd_exportar(
                update, _make_context(["refeicoes", "csv", "2026-08-01", "2026-08-02"])
            )

        # Only the "sem dados" reply, no sync-related messages
        assert update.message.reply_text.await_count == 1
        assert "Sem dados" in update.message.reply_text.call_args[0][0]

    @pytest.mark.asyncio
    async def test_fetches_only_missing_days(self, repo):
        start, end = date(2026, 8, 1), date(2026, 8, 3)
        # Day 1 already has a manual entry -> not "missing"
        repo.save_food_entries(start, [
            {"name": "Ovos", "quantity": 1, "unit": "un", "calories": 100.0, "source": "llm_estimate"},
        ])
        fatsecret = MagicMock()
        fatsecret.get_food_entries.return_value = []
        update = _make_update()
        bot = _make_bot(repo, update.effective_chat.id, fatsecret_client=fatsecret)

        with patch("src.nutrition.fatsecret_mapper.map_fatsecret_entries", return_value=[]), \
             patch("telegram.Bot.send_document", new_callable=AsyncMock):
            await bot._cmd_exportar(update, _make_context(["refeicoes", "csv", str(start), str(end)]))

        called_days = {call.args[0] for call in fatsecret.get_food_entries.call_args_list}
        assert called_days == {date(2026, 8, 2), date(2026, 8, 3)}

    @pytest.mark.asyncio
    async def test_upserts_mapped_entries_for_missing_days(self, repo):
        start = end = date(2026, 8, 5)
        mapped = [
            {"name": "Banana", "calories": 100.0, "protein_g": 1.0, "fat_g": 0.3,
             "carbs_g": 25.0, "fiber_g": 2.0, "quantity": 1.0, "unit": "serving",
             "source": "fatsecret", "barcode": "FS001"},
        ]
        fatsecret = MagicMock()
        fatsecret.get_food_entries.return_value = [{"raw": "entry"}]
        update = _make_update()
        bot = _make_bot(repo, update.effective_chat.id, fatsecret_client=fatsecret)

        with patch("src.nutrition.fatsecret_mapper.map_fatsecret_entries", return_value=mapped), \
             patch("telegram.Bot.send_document", new_callable=AsyncMock):
            await bot._cmd_exportar(update, _make_context(["refeicoes", "csv", str(start), str(end)]))

        entries = repo.get_food_entries(start)
        assert len(entries) == 1
        assert entries[0].name == "Banana"
        assert entries[0].source == "fatsecret"

    @pytest.mark.asyncio
    async def test_sync_failure_does_not_break_export(self, repo):
        """A FatSecret error on the missing day is logged and swallowed; export still runs."""
        start = end = date(2026, 8, 7)
        fatsecret = MagicMock()
        fatsecret.get_food_entries.side_effect = RuntimeError("FatSecret down")
        update = _make_update()
        bot = _make_bot(repo, update.effective_chat.id, fatsecret_client=fatsecret)

        with patch("telegram.Bot.send_document", new_callable=AsyncMock):
            await bot._cmd_exportar(update, _make_context(["refeicoes", "csv", str(start), str(end)]))

        # No crash; export just reports no data for that empty day
        assert any("Sem dados" in c.args[0] for c in update.message.reply_text.call_args_list)

    @pytest.mark.asyncio
    async def test_sync_progress_message_sent_when_days_missing(self, repo):
        start = end = date(2026, 8, 8)
        fatsecret = MagicMock()
        fatsecret.get_food_entries.return_value = []
        update = _make_update()
        bot = _make_bot(repo, update.effective_chat.id, fatsecret_client=fatsecret)

        with patch("src.nutrition.fatsecret_mapper.map_fatsecret_entries", return_value=[]), \
             patch("telegram.Bot.send_document", new_callable=AsyncMock):
            await bot._cmd_exportar(update, _make_context(["refeicoes", "csv", str(start), str(end)]))

        messages = [c.args[0] for c in update.message.reply_text.call_args_list]
        assert any("A sincronizar" in m for m in messages)
