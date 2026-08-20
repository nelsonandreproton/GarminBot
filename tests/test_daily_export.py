"""Tests for /exportar (daily metrics — same data as /sync): the backward-
compatible /exportar N shortcut, and the new [formato] [inicio] [fim] period
export including weight_kg and waist_cm."""

from __future__ import annotations

import json
import os
import tempfile
from datetime import date, timedelta
from unittest.mock import AsyncMock, MagicMock, patch

import pytest

from src.database.repository import Repository
from src.config import Config
from src.telegram.bot import TelegramBot


# ---------------------------------------------------------------------------
# build_csv / build_json / build_xlsx / build_export
# ---------------------------------------------------------------------------

@pytest.fixture
def db_path():
    with tempfile.NamedTemporaryFile(suffix=".db", delete=False) as f:
        path = f.name
    yield path
    try:
        os.unlink(path)
    except PermissionError:
        pass


@pytest.fixture
def repo(db_path):
    r = Repository(db_path)
    r.init_database()
    yield r
    r._engine.dispose()


class TestBuildExport:
    def test_build_csv_contains_expected_columns_and_values(self, repo):
        day = date(2026, 3, 1)
        repo.save_daily_metrics(day, {
            "sleep_hours": 7.5, "sleep_score": 82, "sleep_quality": "Excelente",
            "steps": 10000, "active_calories": 500, "resting_calories": 1700,
            "resting_heart_rate": 55, "avg_stress": 30,
            "body_battery_high": 87, "body_battery_low": 10, "weight_kg": 93.4,
        })
        repo.save_waist_entry(day, 95.5)
        rows = repo.get_metrics_range(day, day)
        waist_by_date = repo.get_waist_range(day, day)
        from src.telegram.daily_export import build_csv
        content = build_csv(rows, waist_by_date).decode("utf-8")
        assert "peso_kg" in content and "perimetro_cm" in content
        assert "93.4" in content
        assert "95.5" in content

    def test_build_json_has_expected_fields(self, repo):
        day = date(2026, 3, 2)
        repo.save_daily_metrics(day, {"steps": 8000, "weight_kg": 90.1})
        repo.save_waist_entry(day, 92.0)
        rows = repo.get_metrics_range(day, day)
        waist_by_date = repo.get_waist_range(day, day)
        from src.telegram.daily_export import build_json
        data = json.loads(build_json(rows, waist_by_date))
        assert data[0]["peso_kg"] == 90.1
        assert data[0]["perimetro_cm"] == 92.0

    def test_build_json_waist_none_when_not_logged(self, repo):
        day = date(2026, 3, 3)
        repo.save_daily_metrics(day, {"steps": 5000})
        rows = repo.get_metrics_range(day, day)
        waist_by_date = repo.get_waist_range(day, day)
        from src.telegram.daily_export import build_json
        data = json.loads(build_json(rows, waist_by_date))
        assert data[0]["perimetro_cm"] is None

    def test_build_xlsx_is_valid_workbook(self, repo):
        day = date(2026, 3, 4)
        repo.save_daily_metrics(day, {"steps": 9000, "weight_kg": 91.0})
        repo.save_waist_entry(day, 93.0)
        rows = repo.get_metrics_range(day, day)
        waist_by_date = repo.get_waist_range(day, day)
        from src.telegram.daily_export import build_xlsx
        content = build_xlsx(rows, waist_by_date)
        import io
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(content))
        rows_out = list(wb.active.iter_rows(values_only=True))
        assert rows_out[1][11] == 91.0  # peso_kg column
        assert rows_out[1][12] == 93.0  # perimetro_cm column

    def test_build_export_dispatches_by_format(self, repo):
        day = date(2026, 3, 5)
        repo.save_daily_metrics(day, {"steps": 1000})
        rows = repo.get_metrics_range(day, day)
        waist_by_date = repo.get_waist_range(day, day)
        from src.telegram.daily_export import build_export
        assert build_export("csv", rows, waist_by_date).startswith(b"data,sono_horas")
        json.loads(build_export("json", rows, waist_by_date))


# ---------------------------------------------------------------------------
# /exportar handler (N shortcut + period/format)
# ---------------------------------------------------------------------------

def _make_config(**overrides):
    defaults = {
        "telegram_bot_token": "fake-token",
        "telegram_chat_id": "123456",
        "garmin_email": "test@example.com",
        "garmin_password": "secret",
        "database_path": ":memory:",
        "log_level": "INFO",
        "log_file": None,
        "daily_alerts": False,
        "groq_api_key": None,
        "usda_api_key": None,
        "api_ninjas_key": None,
        "fatsecret_consumer_key": None,
        "fatsecret_consumer_secret": None,
        "garmin_api_port": None,
        "garmin_api_key": None,
        "health_port": None,
        "newsletter_enabled": False,
    }
    defaults.update(overrides)
    cfg = MagicMock(spec=Config)
    for k, v in defaults.items():
        setattr(cfg, k, v)
    return cfg


_NEXT_CHAT_ID = 1000000


def _make_update():
    global _NEXT_CHAT_ID
    _NEXT_CHAT_ID += 1
    update = MagicMock()
    update.effective_chat.id = _NEXT_CHAT_ID
    update.message = AsyncMock()
    update.message.reply_text = AsyncMock()
    return update


def _make_bot(repo, chat_id):
    cfg = _make_config(telegram_chat_id=str(chat_id))
    bot = TelegramBot(cfg, repo)
    bot._chat_id = chat_id
    return bot


def _make_context(args):
    ctx = MagicMock()
    ctx.args = args
    return ctx


class TestCmdExportarNShortcut:
    """Backward compatibility: /exportar N (bare digit) must keep working exactly
    as before — last N days, CSV, no period/format parsing involved."""

    @pytest.mark.asyncio
    async def test_exportar_n_sends_csv_of_last_n_days(self, repo):
        for i in range(5):
            repo.save_daily_metrics(date.today() - timedelta(days=i), {"steps": 1000 + i})
        update = _make_update()
        bot = _make_bot(repo, update.effective_chat.id)

        with patch("telegram.Bot.send_document", new_callable=AsyncMock) as mock_send:
            await bot._cmd_exportar(update, _make_context(["3"]))

        kwargs = mock_send.call_args.kwargs
        assert kwargs["document"].filename.endswith(".csv")
        assert "3 dias" in kwargs["caption"]

    @pytest.mark.asyncio
    async def test_exportar_no_args_replies_no_data_when_empty(self, repo):
        update = _make_update()
        bot = _make_bot(repo, update.effective_chat.id)
        with patch("telegram.Bot.send_document", new_callable=AsyncMock) as mock_send:
            await bot._cmd_exportar(update, _make_context([]))
        update.message.reply_text.assert_awaited_once()
        assert "Sem dados" in update.message.reply_text.call_args[0][0]
        mock_send.assert_not_awaited()


class TestCmdExportarPeriod:
    @pytest.mark.asyncio
    async def test_invalid_period_replies_with_error(self, repo):
        update = _make_update()
        bot = _make_bot(repo, update.effective_chat.id)
        await bot._cmd_exportar(update, _make_context(["csv", "2026-08-10", "2026-08-01"]))
        update.message.reply_text.assert_awaited_once()
        assert "❌" in update.message.reply_text.call_args[0][0]

    @pytest.mark.asyncio
    async def test_csv_export_includes_weight_and_waist(self, repo):
        today = date.today()
        repo.save_daily_metrics(today, {"steps": 8000, "weight_kg": 93.4})
        repo.save_waist_entry(today, 95.5)
        update = _make_update()
        bot = _make_bot(repo, update.effective_chat.id)

        with patch("telegram.Bot.send_document", new_callable=AsyncMock) as mock_send:
            await bot._cmd_exportar(update, _make_context(["csv", str(today), str(today)]))

        content = mock_send.call_args.kwargs["document"].input_file_content
        assert b"93.4" in content
        assert b"95.5" in content

    @pytest.mark.asyncio
    async def test_json_export_valid(self, repo):
        today = date.today()
        repo.save_daily_metrics(today, {"steps": 8000})
        update = _make_update()
        bot = _make_bot(repo, update.effective_chat.id)

        with patch("telegram.Bot.send_document", new_callable=AsyncMock) as mock_send:
            await bot._cmd_exportar(update, _make_context(["json", str(today), str(today)]))

        content = mock_send.call_args.kwargs["document"].input_file_content
        data = json.loads(content)
        assert data[0]["passos"] == 8000

    @pytest.mark.asyncio
    async def test_xlsx_export_is_valid_workbook(self, repo):
        today = date.today()
        repo.save_daily_metrics(today, {"steps": 8000})
        update = _make_update()
        bot = _make_bot(repo, update.effective_chat.id)

        with patch("telegram.Bot.send_document", new_callable=AsyncMock) as mock_send:
            await bot._cmd_exportar(update, _make_context(["xlsx", str(today), str(today)]))

        content = mock_send.call_args.kwargs["document"].input_file_content
        import io
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(content))
        assert wb.active.iter_rows(values_only=True)

    @pytest.mark.asyncio
    async def test_custom_period_passed_to_repository(self, repo):
        update = _make_update()
        bot = _make_bot(repo, update.effective_chat.id)

        with patch.object(repo, "get_metrics_range", wraps=repo.get_metrics_range) as mock_range:
            await bot._cmd_exportar(
                update, _make_context(["csv", "2026-08-01", "2026-08-10"])
            )

        mock_range.assert_called_once_with(date(2026, 8, 1), date(2026, 8, 10))

    @pytest.mark.asyncio
    async def test_no_data_in_period_replies_with_message(self, repo):
        update = _make_update()
        bot = _make_bot(repo, update.effective_chat.id)
        with patch("telegram.Bot.send_document", new_callable=AsyncMock) as mock_send:
            await bot._cmd_exportar(update, _make_context(["csv", "2020-01-01", "2020-01-02"]))
        assert "Sem dados" in update.message.reply_text.call_args[0][0]
        mock_send.assert_not_awaited()

    @pytest.mark.asyncio
    async def test_defaults_to_csv_last_90_days_with_no_args(self, repo):
        repo.save_daily_metrics(date.today(), {"steps": 100})
        update = _make_update()
        bot = _make_bot(repo, update.effective_chat.id)

        with patch.object(repo, "get_metrics_range", wraps=repo.get_metrics_range) as mock_range, \
             patch("telegram.Bot.send_document", new_callable=AsyncMock):
            await bot._cmd_exportar(update, _make_context([]))

        called_start, called_end = mock_range.call_args.args
        assert called_end == date.today()
        assert called_start == date.today() - timedelta(days=90)
