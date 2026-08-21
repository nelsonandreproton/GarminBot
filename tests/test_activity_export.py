"""Tests for /exportar treinos: CSV/JSON/XLSX export of garmin_activities."""

from __future__ import annotations

import json
import os
import tempfile
from datetime import date, timedelta
from unittest.mock import AsyncMock, MagicMock, patch

import pytest

from src.database.repository import Repository
from src.config import Config
from src.telegram.bot import TelegramBot


# ---------------------------------------------------------------------------
# build_csv / build_json / build_xlsx / build_export
# ---------------------------------------------------------------------------

@pytest.fixture
def db_path():
    with tempfile.NamedTemporaryFile(suffix=".db", delete=False) as f:
        path = f.name
    yield path
    try:
        os.unlink(path)
    except PermissionError:
        pass


@pytest.fixture
def repo(db_path):
    r = Repository(db_path)
    r.init_database()
    yield r
    r._engine.dispose()


class TestBuildExport:
    def test_build_csv_contains_expected_columns_and_values(self, repo):
        day = date(2026, 3, 1)
        repo.upsert_garmin_activity(
            1, day, "Caminhada", "walking", 40, 250, 3.5,
            avg_hr=110, max_hr=140, min_hr=88,
        )
        entries = repo.get_garmin_activities_range(day, day)
        from src.telegram.activity_export import build_csv
        content = build_csv(entries).decode("utf-8")
        assert "data,tipo,duracao_min,calorias,bpm_min,bpm_medio,bpm_max,distancia_km,nome" in content
        assert "walking" in content
        assert "88" in content
        assert "Caminhada" in content

    def test_build_json_has_expected_fields(self, repo):
        day = date(2026, 3, 2)
        repo.upsert_garmin_activity(
            2, day, "Gym", "strength_training", 45, 320, None,
            avg_hr=128, max_hr=165, min_hr=95,
        )
        entries = repo.get_garmin_activities_range(day, day)
        from src.telegram.activity_export import build_json
        data = json.loads(build_json(entries))
        assert data[0]["tipo"] == "strength_training"
        assert data[0]["bpm_min"] == 95
        assert data[0]["bpm_medio"] == 128
        assert data[0]["bpm_max"] == 165
        assert data[0]["nome"] == "Gym"

    def test_build_xlsx_is_valid_workbook(self, repo):
        day = date(2026, 3, 3)
        repo.upsert_garmin_activity(3, day, "Corrida", "running", 30, 280, 5.0, avg_hr=150, max_hr=175, min_hr=100)
        entries = repo.get_garmin_activities_range(day, day)
        from src.telegram.activity_export import build_xlsx
        content = build_xlsx(entries)
        import io
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(content))
        rows = list(wb.active.iter_rows(values_only=True))
        assert rows[1][1] == "running"
        assert rows[1][8] == "Corrida"

    def test_build_export_dispatches_by_format(self, repo):
        day = date(2026, 3, 4)
        repo.upsert_garmin_activity(4, day, "Walk", "walking", 20, 100, None)
        entries = repo.get_garmin_activities_range(day, day)
        from src.telegram.activity_export import build_export
        assert build_export("csv", entries).startswith(b"data,tipo")
        json.loads(build_export("json", entries))


# ---------------------------------------------------------------------------
# /exportar treinos handler
# ---------------------------------------------------------------------------

def _make_config(**overrides):
    defaults = {
        "telegram_bot_token": "fake-token",
        "telegram_chat_id": "123456",
        "garmin_email": "test@example.com",
        "garmin_password": "secret",
        "database_path": ":memory:",
        "log_level": "INFO",
        "log_file": None,
        "daily_alerts": False,
        "groq_api_key": None,
        "fatsecret_consumer_key": None,
        "fatsecret_consumer_secret": None,
        "garmin_api_port": None,
        "garmin_api_key": None,
        "health_port": None,
        "newsletter_enabled": False,
    }
    defaults.update(overrides)
    cfg = MagicMock(spec=Config)
    for k, v in defaults.items():
        setattr(cfg, k, v)
    return cfg


_NEXT_CHAT_ID = 900000


def _make_update():
    global _NEXT_CHAT_ID
    _NEXT_CHAT_ID += 1
    update = MagicMock()
    update.effective_chat.id = _NEXT_CHAT_ID
    update.message = AsyncMock()
    update.message.reply_text = AsyncMock()
    return update


def _make_bot(repo, chat_id):
    cfg = _make_config(telegram_chat_id=str(chat_id))
    bot = TelegramBot(cfg, repo)
    bot._chat_id = chat_id
    return bot


def _make_context(args):
    ctx = MagicMock()
    ctx.args = args
    return ctx


class TestCmdExportarTreinos:
    @pytest.mark.asyncio
    async def test_no_activities_replies_with_message_and_sends_nothing(self, repo):
        update = _make_update()
        bot = _make_bot(repo, update.effective_chat.id)
        with patch("telegram.Bot.send_document", new_callable=AsyncMock) as mock_send:
            await bot._cmd_exportar(update, _make_context(["treinos"]))
        update.message.reply_text.assert_awaited_once()
        assert "Sem dados de treinos" in update.message.reply_text.call_args[0][0]
        mock_send.assert_not_awaited()

    @pytest.mark.asyncio
    async def test_invalid_period_replies_with_error(self, repo):
        update = _make_update()
        bot = _make_bot(repo, update.effective_chat.id)
        await bot._cmd_exportar(update, _make_context(["treinos", "csv", "2026-08-10", "2026-08-01"]))
        update.message.reply_text.assert_awaited_once()
        assert "❌" in update.message.reply_text.call_args[0][0]

    @pytest.mark.asyncio
    async def test_csv_export_sends_document(self, repo):
        today = date.today()
        repo.upsert_garmin_activity(10, today, "Caminhada", "walking", 40, 250, 3.5, avg_hr=110, max_hr=140, min_hr=88)
        update = _make_update()
        bot = _make_bot(repo, update.effective_chat.id)

        with patch("telegram.Bot.send_document", new_callable=AsyncMock) as mock_send:
            await bot._cmd_exportar(update, _make_context(["treinos"]))

        mock_send.assert_awaited_once()
        kwargs = mock_send.call_args.kwargs
        assert kwargs["chat_id"] == bot._chat_id
        assert kwargs["document"].filename.endswith(".csv")
        content = kwargs["document"].input_file_content
        assert b"Caminhada" in content

    @pytest.mark.asyncio
    async def test_json_export_contains_valid_json(self, repo):
        today = date.today()
        repo.upsert_garmin_activity(11, today, "Gym", "strength_training", 45, 320, None, avg_hr=128, max_hr=165, min_hr=95)
        update = _make_update()
        bot = _make_bot(repo, update.effective_chat.id)

        with patch("telegram.Bot.send_document", new_callable=AsyncMock) as mock_send:
            await bot._cmd_exportar(update, _make_context(["treinos", "json"]))

        content = mock_send.call_args.kwargs["document"].input_file_content
        data = json.loads(content)
        assert data[0]["nome"] == "Gym"
        assert data[0]["bpm_min"] == 95

    @pytest.mark.asyncio
    async def test_xlsx_export_is_valid_workbook(self, repo):
        today = date.today()
        repo.upsert_garmin_activity(12, today, "Corrida", "running", 30, 280, 5.0, avg_hr=150, max_hr=175, min_hr=100)
        update = _make_update()
        bot = _make_bot(repo, update.effective_chat.id)

        with patch("telegram.Bot.send_document", new_callable=AsyncMock) as mock_send:
            await bot._cmd_exportar(update, _make_context(["treinos", "xlsx"]))

        content = mock_send.call_args.kwargs["document"].input_file_content
        import io
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(content))
        rows = list(wb.active.iter_rows(values_only=True))
        assert rows[1][8] == "Corrida"

    @pytest.mark.asyncio
    async def test_custom_period_passed_to_repository(self, repo):
        update = _make_update()
        bot = _make_bot(repo, update.effective_chat.id)

        with patch.object(repo, "get_garmin_activities_range", wraps=repo.get_garmin_activities_range) as mock_range:
            await bot._cmd_exportar(
                update, _make_context(["treinos", "csv", "2026-08-01", "2026-08-10"])
            )

        mock_range.assert_called_once_with(date(2026, 8, 1), date(2026, 8, 10))

    @pytest.mark.asyncio
    async def test_caption_mentions_count_and_period(self, repo):
        today = date.today()
        repo.upsert_garmin_activity(13, today, "Walk", "walking", 20, 100, None)
        update = _make_update()
        bot = _make_bot(repo, update.effective_chat.id)

        with patch("telegram.Bot.send_document", new_callable=AsyncMock) as mock_send:
            await bot._cmd_exportar(update, _make_context(["treinos"]))

        caption = mock_send.call_args.kwargs["caption"]
        assert "1 treinos" in caption
